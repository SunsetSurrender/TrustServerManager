#!/usr/bin/env python3
"""Seriali VERI degli armadi: dal foglio «Seriali rack» a deploy/baseline/rack-serials.json.

Il rilievo si fa in sala: si riconosce il rack dall'etichetta sulla pianta
(oggi spesso il codice, «R01»), si legge la targhetta, e si scrive il numero
nelle colonne «seriale effettivo» del modello Excel. Questo script trasforma
quelle righe nel file di correzioni che `tools/make-production-baseline.py`
applica sopra il seed.

Regole di una riga:
- entrambe le colonne «seriale effettivo» VUOTE  → il rack non cambia;
- «seriale effettivo 1» = NESSUNO                → il rack resta senza seriali;
- altrimenti                                     → i seriali diventano esattamente
  quelli scritti (1 o 2), al posto di quelli attuali.

La riga si riconosce per `_uid`, non per codice: il codice si può rinominare,
l'`_uid` no. Il codice serve solo da controllo incrociato.

Il file prodotto si UNISCE a quello esistente: un secondo giro di rilievo (una
sala alla volta) aggiunge, non cancella il primo.

Uso:
    python tools/apply-rack-serials.py "compilazione/Modello dispositivi sale.xlsx" [--dry-run]
poi:
    python tools/make-production-baseline.py            # rigenera e mostra il digest
    (leggere deploy/baseline/MANIFEST.md, sezione «Seriali … corretti»)
    python tools/make-production-baseline.py --aggiorna-digest
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BASELINE = ROOT / "deploy" / "baseline" / "production-baseline.json"
CORREZIONI = ROOT / "deploy" / "baseline" / "rack-serials.json"
FOGLIO = "Seriali rack"
NESSUNO = "NESSUNO"
UID = re.compile(r"^[0-9a-f]{8}-[0-9a-f]{4}-4[0-9a-f]{3}-[89ab][0-9a-f]{3}-[0-9a-f]{12}$")
sys.stdout.reconfigure(encoding="utf-8")


class Errore(Exception):
    pass


def _testo(v) -> str:
    """Cella → testo. Un numero intero letto come float (2020062936.0) torna intero:
    succede se la colonna ha perso il formato Testo."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def leggi(xlsx: Path) -> list[dict]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if FOGLIO not in wb.sheetnames:
        raise Errore(f"{xlsx}: manca il foglio «{FOGLIO}»")
    righe = list(wb[FOGLIO].iter_rows(values_only=True))
    if not righe:
        raise Errore(f"«{FOGLIO}» è vuoto")
    intest = [_testo(h).lower() for h in righe[0]]
    def col(nome: str) -> int:
        if nome not in intest:
            raise Errore(f"«{FOGLIO}»: manca la colonna «{nome}»")
        return intest.index(nome)
    c_uid, c_rack = col("_uid"), col("rack")
    c_s1, c_s2 = col("seriale effettivo 1"), col("seriale effettivo 2")
    out = []
    for n, r in enumerate(righe[1:], start=2):
        r = list(r) + [None] * len(intest)
        out.append({"riga": n, "_uid": _testo(r[c_uid]), "rack": _testo(r[c_rack]),
                    "s1": _testo(r[c_s1]), "s2": _testo(r[c_s2])})
    return out


def rack_della_baseline() -> dict[str, dict]:
    doc = json.loads(BASELINE.read_text(encoding="utf-8"))
    return {K["_uid"]: {"rack": K["id"], "sala": S["id"], "seriali": list(K.get("seriali") or [])}
            for L in doc["locations"] for S in L["sale"] for K in S["racks"]}


def correzioni_da(righe: list[dict], rack: dict[str, dict]) -> tuple[dict, list[str]]:
    nuove: dict[str, dict] = {}
    errori: list[str] = []
    visti_seriali: dict[str, str] = {}
    for r in righe:
        if not (r["s1"] or r["s2"]):
            continue
        dove = f"riga {r['riga']} ({r['rack'] or '?'})"
        if not UID.match(r["_uid"]):
            errori.append(f"{dove}: _uid mancante o alterato — non modificare le prime colonne")
            continue
        if r["_uid"] not in rack:
            errori.append(f"{dove}: _uid {r['_uid']} non è un rack della baseline")
            continue
        if r["rack"] != rack[r["_uid"]]["rack"]:
            errori.append(f"{dove}: il codice non combacia con l'_uid "
                          f"(atteso {rack[r['_uid']]['rack']!r}) — righe spostate?")
            continue
        if r["s1"].upper() == NESSUNO:
            if r["s2"]:
                errori.append(f"{dove}: {NESSUNO} nella colonna 1 e un seriale nella 2")
                continue
            seriali = []
        elif not r["s1"]:
            errori.append(f"{dove}: seriale effettivo 2 senza il seriale 1")
            continue
        else:
            seriali = [s for s in (r["s1"], r["s2"]) if s]
            if len(set(seriali)) != len(seriali):
                errori.append(f"{dove}: lo stesso seriale due volte")
                continue
        for s in seriali:
            if s in visti_seriali and visti_seriali[s] != r["rack"]:
                errori.append(f"{dove}: il seriale {s!r} è già scritto sul rack {visti_seriali[s]}")
            visti_seriali[s] = r["rack"]
        if r["_uid"] in nuove:
            errori.append(f"{dove}: rack ripetuto nel foglio")
            continue
        nuove[r["_uid"]] = {"rack": r["rack"], "sala": rack[r["_uid"]]["sala"], "seriali": seriali}
    return nuove, errori


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("xlsx", type=Path)
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    try:
        righe = leggi(args.xlsx)
    except Errore as exc:
        print(f"FERMO: {exc}", file=sys.stderr)
        return 2
    rack = rack_della_baseline()
    nuove, errori = correzioni_da(righe, rack)
    if errori:
        print("FERMO: il foglio ha righe non valide, niente scritto:", file=sys.stderr)
        for e in errori:
            print(f"  - {e}", file=sys.stderr)
        return 1

    esistenti = {}
    if CORREZIONI.exists():
        esistenti = json.loads(CORREZIONI.read_text(encoding="utf-8")).get("rack", {})
    unite = {**esistenti, **nuove}

    print(f"righe con un seriale effettivo: {len(nuove)}  (correzioni totali nel file: {len(unite)})")
    for uid, v in nuove.items():
        prima = rack[uid]["seriali"]
        segno = "=" if prima == v["seriali"] else "→"
        print(f"  {v['sala']:<14} {v['rack']:<10} {', '.join(prima) or '—'}  {segno}  {', '.join(v['seriali']) or '—'}")
    if args.dry_run:
        print("--dry-run: niente scritto")
        return 0
    # Ordinato per sala e codice: il file si rilegge in una revisione.
    ordine = sorted(unite.items(), key=lambda kv: (kv[1]["sala"], kv[1]["rack"]))
    testo = json.dumps({"schema": 1, "rack": dict(ordine)}, ensure_ascii=False, indent=2) + "\n"
    CORREZIONI.write_text(testo, encoding="utf-8", newline="\n")
    print(f"\nscritto {CORREZIONI.relative_to(ROOT)}")
    print("ora: python tools/make-production-baseline.py   (poi leggere il manifesto)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
