#!/usr/bin/env python3
"""Modello Excel per compilare a mano i dispositivi delle sale.

Il file prodotto è nel formato dell'import dell'applicazione (pannello «Importa
dati», ↑ CSV/XLSX): l'import legge SOLO il PRIMO foglio, con le intestazioni
`_uid, location, sala, rack, nome, tipo, stato, presenza, modello, ip, seriale,
referente, u, h, garanzia, supporto, note`.

La struttura (siti, sale, rack) si legge dalla BASELINE DI PRODUZIONE, non dal
seed di sviluppo: le tendine offrono esattamente i rack che esisteranno
sull'host dopo `production_baseline.py --init`. Se la baseline cambia, il
modello si rigenera.

Scelte che contano per l'import:
- `location` è una FORMULA ricavata dalla sala (gli id delle sale sono unici):
  una colonna in meno da sbagliare. Excel salva il valore calcolato insieme
  alla formula, ed è quel valore che l'import legge.
- `rack` è una tendina DIPENDENTE dalla sala scelta.
- colonne di testo in formato Testo (`@`): senza, Excel trasforma
  `2027-12-31` in un numero di serie di data e un seriale `0012345` in `12345`.
- nessuna riga pre-compilata: una riga di esempio dimenticata verrebbe
  importata come un dispositivo vero.

Uso:
    python tools/make-import-template.py [--baseline deploy/baseline/production-baseline.json]
                                         [--out "compilazione/Modello dispositivi sale.xlsx"]
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = ["_uid", "location", "sala", "rack", "nome", "tipo", "stato", "presenza",
           "modello", "ip", "seriale", "referente", "u", "h", "garanzia", "supporto", "note"]
TIPI = ["server", "rete", "storage", "firewall", "alimentazione", "altro"]
STATI = ["attivo", "manutenzione", "dismissione", "dismesso"]
PRESENZE = ["presente", "rimosso"]
RIGHE = 1500          # righe con tendine e formule già pronte

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
AUTO_FILL = PatternFill("solid", fgColor="EDEDED")
WARN_FILL = PatternFill("solid", fgColor="FFC7CE")
WIDTH = {"_uid": 10, "location": 15, "sala": 16, "rack": 12, "nome": 24, "tipo": 14,
         "stato": 14, "presenza": 11, "modello": 22, "ip": 16, "seriale": 18,
         "referente": 18, "u": 6, "h": 6, "garanzia": 12, "supporto": 12, "note": 40}


def col(name: str) -> str:
    return chr(ord("A") + HEADERS.index(name))


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--baseline", default="deploy/baseline/production-baseline.json")
    ap.add_argument("--out", default="compilazione/Modello dispositivi sale.xlsx")
    args = ap.parse_args()

    doc = json.loads(Path(args.baseline).read_text(encoding="utf-8"))
    sale, rack, seriali_rack = [], [], []
    for L in doc["locations"]:
        for S in L["sale"]:
            sale.append((S["id"], S.get("nome", S["id"]), L["id"], L.get("nome", L["id"])))
            for K in S["racks"]:
                rack.append((K["id"], K.get("name", ""), K.get("row", "") or "", int(K.get("u") or 0),
                             S["id"], L["id"]))
                seriali_rack.append((K["_uid"], L["id"], S["id"], K["id"],
                                     (K.get("seriali") or [K["id"]])[0], K.get("name", ""),
                                     K.get("row", "") or "", int(K.get("u") or 0),
                                     " / ".join(K.get("seriali") or [])))
    ids_sale = [s[0] for s in sale]
    assert len(set(ids_sale)) == len(ids_sale), "id di sala ripetuti: la formula della location non vale"
    ids_rack = [r[0] for r in rack]
    assert len(set(ids_rack)) == len(ids_rack), "id di rack ripetuti fra sale"

    wb = Workbook()
    # Excel ricalcola all'apertura: le formule di `location` hanno un valore salvato.
    wb.calculation.fullCalcOnLoad = True

    # ---- 1. Dispositivi: DEVE restare il primo foglio --------------------
    ws = wb.active
    ws.title = "Dispositivi"
    ws.append(HEADERS)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for h in HEADERS:
        cd = ws.column_dimensions[col(h)]
        cd.width = WIDTH[h]
        if h not in ("u", "h"):
            cd.number_format = "@"

    last = RIGHE + 1
    L, S, K, U, H = col("location"), col("sala"), col("rack"), col("u"), col("h")
    for r in range(2, last + 1):
        ws[f"{L}{r}"] = f'=IF(${S}{r}="","",IFERROR(VLOOKUP(${S}{r},Sale!$A:$C,3,FALSE),"sala?"))'
        ws[f"{L}{r}"].number_format = "General"

    def lista(valori, cella, titolo, msg):
        dv = DataValidation(type="list", formula1='"' + ",".join(valori) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle=titolo, error=msg)
        ws.add_data_validation(dv)
        dv.add(f"{cella}2:{cella}{last}")

    n_sale = len(sale)
    dv = DataValidation(type="list", formula1=f"=Sale!$A$2:$A${n_sale + 1}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Sala", error="Scegli una sala dall'elenco.")
    ws.add_data_validation(dv)
    dv.add(f"{S}2:{S}{last}")
    # Tendina dei rack della sala scelta: il foglio Rack è ordinato per sala.
    dv = DataValidation(
        type="list", allow_blank=True, showErrorMessage=True, errorTitle="Rack",
        error="Scegli prima la sala, poi un rack di quella sala.",
        formula1=f"=OFFSET(Rack!$A$1,MATCH(${S}2,Rack!$E:$E,0)-1,0,COUNTIF(Rack!$E:$E,${S}2),1)")
    ws.add_data_validation(dv)
    dv.add(f"{K}2:{K}{last}")
    lista(TIPI, col("tipo"), "Tipo", "Valori ammessi: " + ", ".join(TIPI))
    lista(STATI, col("stato"), "Stato", "Valori ammessi: " + ", ".join(STATI))
    lista(PRESENZE, col("presenza"), "Presenza", "presente o rimosso")
    for c, nome in ((U, "u"), (H, "h")):
        dv = DataValidation(type="whole", operator="between", formula1="1", formula2="60",
                            allow_blank=True, showErrorMessage=True, errorTitle=nome,
                            error="Numero intero di unità rack.")
        ws.add_data_validation(dv)
        dv.add(f"{c}2:{c}{last}")
    for c in (col("garanzia"), col("supporto")):
        dv = DataValidation(type="textLength", operator="equal", formula1="10", allow_blank=True,
                            showErrorMessage=True, showInputMessage=True,
                            promptTitle="Data", prompt="Formato AAAA-MM-GG, es. 2027-12-31",
                            errorTitle="Data", error="Scrivi la data come AAAA-MM-GG (10 caratteri).")
        ws.add_data_validation(dv)
        dv.add(f"{c}2:{c}{last}")
    dv = DataValidation(type="textLength", operator="equal", formula1="0", allow_blank=True,
                        showErrorMessage=True, showInputMessage=True, promptTitle="_uid",
                        prompt="Lasciare VUOTO: lo assegna l'applicazione all'import.",
                        errorTitle="_uid", error="Per i dispositivi nuovi _uid resta vuoto.")
    ws.add_data_validation(dv)
    dv.add(f"A2:A{last}")

    # Evidenziazioni: colonne calcolate in grigio, dispositivo che esce dal rack in rosso.
    ws.conditional_formatting.add(f"A2:A{last}", FormulaRule(formula=["TRUE"], fill=AUTO_FILL))
    ws.conditional_formatting.add(f"{L}2:{L}{last}", FormulaRule(formula=["TRUE"], fill=AUTO_FILL))
    ws.conditional_formatting.add(
        f"{U}2:{H}{last}",
        FormulaRule(formula=[f'AND(${K}2<>"",${U}2<>"",${U}2+MAX(1,N(${H}2))-1>IFERROR(VLOOKUP(${K}2,Rack!$A:$D,4,FALSE),999))'],
                    fill=WARN_FILL))
    ws.conditional_formatting.add(
        f"{col('nome')}2:{col('nome')}{last}",
        FormulaRule(formula=[f'AND(${K}2<>"",TRIM(${col("nome")}2)="")'], fill=WARN_FILL))
    ws.auto_filter.ref = f"A1:{col('note')}{last}"

    # ---- 2. Seriali rack: il rilievo delle targhette degli armadi ---------
    # NON si importa dall'applicazione: lo legge tools/apply-rack-serials.py
    # PRIMA del deploy, e i seriali entrano nella baseline. Le colonne fino a
    # «seriali attuali» sono riferimento e non si toccano; si compilano solo le
    # due «seriale effettivo».
    wsr = wb.create_sheet("Seriali rack")
    intest = ["_uid", "location", "sala", "rack", "etichetta sulla pianta", "nome rack", "fila",
              "U", "seriali attuali", "seriale effettivo 1", "seriale effettivo 2", "note"]
    wsr.append(intest)
    for c in wsr[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
    for c in wsr[1][9:11]:
        c.fill = PatternFill("solid", fgColor="BF8F00")
    for r in seriali_rack:
        wsr.append(list(r) + ["", "", ""])
    n = len(seriali_rack) + 1
    for col_, w in zip("ABCDEFGHIJKL", (10, 14, 15, 10, 16, 28, 6, 5, 22, 20, 20, 30)):
        wsr.column_dimensions[col_].width = w
    for col_ in ("J", "K", "L"):
        wsr.column_dimensions[col_].number_format = "@"
        for row in range(2, n + 1):
            wsr[f"{col_}{row}"].number_format = "@"
            if col_ != "L":
                wsr[f"{col_}{row}"].fill = INPUT_FILL
    wsr.column_dimensions["A"].hidden = True       # _uid: serve allo script, non a chi compila
    wsr.freeze_panes = "E2"
    wsr.auto_filter.ref = f"A1:L{n}"
    dv = DataValidation(type="custom", formula1='=AND(J2<>"",J2<>K2)', allow_blank=True,
                        showErrorMessage=True, errorTitle="Seriale 2",
                        error="Prima il seriale 1, e diverso dal 2.")
    wsr.add_data_validation(dv)
    dv.add(f"K2:K{n}")
    dv = DataValidation(type="custom", formula1="=TRUE", allow_blank=True, showInputMessage=True,
                        promptTitle="Seriale effettivo",
                        prompt="Il numero sulla targhetta dell'armadio. Vuoto = non cambia. "
                               "NESSUNO = il rack resta senza seriale.")
    wsr.add_data_validation(dv)
    dv.add(f"J2:J{n}")

    # ---- 3. Istruzioni ----------------------------------------------------
    wi = wb.create_sheet("Istruzioni")
    istr = [
        ("Come compilare", ""),
        ("1", "Una riga per dispositivo, nel foglio «Dispositivi». Non riordinare i fogli: l'import legge SOLO il primo."),
        ("2", "Scegli la SALA dalla tendina, poi il RACK (la tendina mostra solo i rack di quella sala). La location si compila da sola."),
        ("3", "Obbligatori: sala, rack, nome. Consigliati: tipo, stato, u, h."),
        ("4", "_uid: lasciare VUOTO. Lo assegna l'applicazione."),
        ("5", "u = unità di PARTENZA contata dal basso; h = altezza in unità (default 1). In rosso se il dispositivo esce dal rack."),
        ("6", "Date (garanzia, supporto) come testo AAAA-MM-GG, es. 2027-12-31."),
        ("7", "Il nome NON è identità, ma due dispositivi con lo stesso nome nello stesso rack rendono la riga ambigua: usare nomi distinti."),
        ("Importare", ""),
        ("a", "Nell'applicazione: modalità Editing (Ctrl+E) → Importa dati → scegliere questo file."),
        ("b", "L'anteprima elenca nuovi / aggiornati / saltati con il motivo. Applicare solo se l'anteprima è quella attesa."),
        ("c", "Si può importare a più riprese (es. una sala alla volta): un dispositivo già importato si aggiorna dall'export, che porta il suo _uid."),
        ("Seriali rack", ""),
        ("1", "Foglio «Seriali rack»: si compila PRIMA del deploy. Riconosci il rack dall'«etichetta sulla pianta» (quella che vedi nel disegno della sala) e scrivi il seriale della targhetta in «seriale effettivo 1» (e 2 se ce ne sono due)."),
        ("2", "Vuoto = il rack non cambia. NESSUNO = il rack resta senza seriale. Non modificare le colonne grigie e non riordinare le righe."),
        ("3", "Quel foglio NON si importa dall'applicazione: lo legge tools/apply-rack-serials.py, e i seriali entrano nella struttura iniziale. Dopo il deploy le correzioni si fanno dalla scheda del rack (⚙ Rack)."),
        ("Valori ammessi", ""),
        ("tipo", ", ".join(TIPI)),
        ("stato", ", ".join(STATI) + " — ciclo di vita operativo"),
        ("presenza", "presente, rimosso — se l'hardware occupa ancora lo slot. Vuoto = presente"),
    ]
    for a, b in istr:
        wi.append([a, b])
        if not b:
            wi.cell(wi.max_row, 1).font = Font(bold=True, size=12)
    wi.column_dimensions["A"].width = 16
    wi.column_dimensions["B"].width = 120

    # ---- 4. Rack (ordinato per sala: la tendina dipendente ci conta) ------
    wr = wb.create_sheet("Rack")
    wr.append(["rack", "nome rack", "fila", "unità", "sala", "location"])
    for c in wr[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for r in sorted(rack, key=lambda x: (ids_sale.index(x[4]), ids_rack.index(x[0]))):
        wr.append(list(r))
    for c, w in zip("ABCDEF", (12, 30, 6, 7, 16, 15)):
        wr.column_dimensions[c].width = w
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:F{len(rack) + 1}"

    # ---- 5. Sale -----------------------------------------------------------
    wsl = wb.create_sheet("Sale")
    wsl.append(["sala", "nome sala", "location", "nome location", "rack"])
    for c in wsl[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    for s in sale:
        wsl.append([s[0], s[1], s[2], s[3], sum(1 for r in rack if r[4] == s[0])])
    for c, w in zip("ABCDE", (16, 18, 15, 18, 6)):
        wsl.column_dimensions[c].width = w

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{out}: {len(sale)} sale, {len(rack)} rack, {RIGHE} righe pronte")


if __name__ == "__main__":
    main()
